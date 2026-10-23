#!/usr/bin/env python3
"""Convert Markdown report to a beautifully formatted Word document (.docx)."""
from __future__ import annotations

import re
import sys
from pathlib import Path

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_formatted_run(p, text):
    """Parse inline **bold** and `code` markers."""
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = p.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('`') and token.endswith('`'):
            run = p.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(160, 40, 40)
        else:
            p.add_run(token)


def markdown_to_docx(md_path: Path, docx_path: Path):
    doc = docx.Document()

    # Set page margins (1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    lines = md_path.read_text(encoding='utf-8').splitlines()

    in_code_block = False
    code_block_lang = ""
    code_lines = []

    in_table = False
    table_lines = []

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            return
        
        # Parse table rows
        rows_data = []
        for l in table_lines:
            if re.match(r'^\s*\|?\s*[-:]+[-| :]*\|?\s*$', l):
                continue  # separator
            cols = [c.strip() for c in l.strip().strip('|').split('|')]
            rows_data.append(cols)

        if rows_data:
            num_cols = max(len(r) for r in rows_data)
            # Normalize column counts
            for r in rows_data:
                while len(r) < num_cols:
                    r.append("")

            tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

            for r_idx, row in enumerate(rows_data):
                for c_idx, val in enumerate(row):
                    cell = tbl.cell(r_idx, c_idx)
                    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.1

                    if r_idx == 0:
                        set_cell_background(cell, '1F497D')
                        run = p.add_run(val)
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9.5)
                    else:
                        bg_color = 'F2F2F2' if r_idx % 2 == 1 else 'FFFFFF'
                        set_cell_background(cell, bg_color)
                        add_formatted_run(p, val)
                        for r in p.runs:
                            if not r.font.size:
                                r.font.size = Pt(9.0)

            doc.add_paragraph()  # spacer

        table_lines = []
        in_table = False

    def flush_code_block():
        nonlocal in_code_block, code_block_lang, code_lines
        if not code_lines:
            in_code_block = False
            return
        
        code_text = "\n".join(code_lines)
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, 'F5F5F5')
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(40, 40, 40)

        doc.add_paragraph()  # spacer
        code_lines = []
        in_code_block = False

    i = 0
    while i < len(lines):
        line = lines[i]

        # Check code block fences
        if line.strip().startswith('```'):
            if not in_code_block:
                if in_table:
                    flush_table()
                in_code_block = True
                code_block_lang = line.strip()[3:].strip()
                code_lines = []
            else:
                flush_code_block()
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Check table lines
        if line.strip().startswith('|') and line.strip().endswith('|'):
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            flush_table()

        # Headings
        if line.startswith('# '):
            h = doc.add_heading(line[2:].strip(), level=1)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(6)
            for r in h.runs:
                r.font.color.rgb = RGBColor(31, 73, 125)
        elif line.startswith('## '):
            h = doc.add_heading(line[3:].strip(), level=2)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(5)
            for r in h.runs:
                r.font.color.rgb = RGBColor(54, 95, 145)
        elif line.startswith('### '):
            h = doc.add_heading(line[4:].strip(), level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            for r in h.runs:
                r.font.color.rgb = RGBColor(79, 129, 189)
        elif line.startswith('#### '):
            h = doc.add_heading(line[5:].strip(), level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(3)
        # Blockquote / Alert
        elif line.strip().startswith('>'):
            content = line.strip().lstrip('>').strip()
            tbl = doc.add_table(rows=1, cols=1)
            cell = tbl.cell(0, 0)
            set_cell_background(cell, 'EBF1F5')
            set_cell_margins(cell, top=100, bottom=100, left=180, right=180)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_run(p, content)
            for r in p.runs:
                r.font.size = Pt(9.5)
                r.font.italic = True
        # Horizontal rule
        elif line.strip() in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p_border = p.paragraph_format
        # Unordered list
        elif line.strip().startswith('- ') or line.strip().startswith('* '):
            item_text = line.strip()[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_run(p, item_text)
        # Ordered list
        elif re.match(r'^\d+\.\s', line.strip()):
            item_text = re.sub(r'^\d+\.\s', '', line.strip())
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_run(p, item_text)
        # Empty line
        elif not line.strip():
            pass
        # Normal paragraph
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_run(p, line.strip())

        i += 1

    if in_table:
        flush_table()
    if in_code_block:
        flush_code_block()

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    print(f"Successfully generated DOCX report: {docx_path} ({docx_path.stat().st_size} bytes)")


if __name__ == '__main__':
    md_file = ROOT / "artifacts/track1_final/reports/SYSTEM_ARCHITECTURE_AND_DEPLOYMENT_REPORT.md"
    docx_file = ROOT / "artifacts/track1_final/reports/SYSTEM_ARCHITECTURE_AND_DEPLOYMENT_REPORT.docx"
    markdown_to_docx(md_file, docx_file)
